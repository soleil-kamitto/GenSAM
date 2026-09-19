"""
Convierte docs/capitulo1.md a docs/capitulo1.docx con formato Word limpio.
Uso: python scripts/md_to_docx.py
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

MD_PATH   = Path('docs/capitulo1.md')
OUT_PATH  = Path('docs/capitulo1.docx')
BASE_DIR  = MD_PATH.parent  # las rutas de imagen en el markdown son relativas aqui

doc = Document()

# ── estilos base ──────────────────────────────────────────────────────────────
style_normal = doc.styles['Normal']
style_normal.font.name = 'Times New Roman'
style_normal.font.size = Pt(12)

for h_name in ['Heading 1', 'Heading 2', 'Heading 3']:
    s = doc.styles[h_name]
    s.font.name = 'Times New Roman'
    s.font.color.rgb = RGBColor(0, 0, 0)

# margenes
for section in doc.sections:
    section.top_margin    = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin   = Inches(1.18)
    section.right_margin  = Inches(1.18)


def add_paragraph(doc, text, bold_parts=None, style='Normal'):
    """Agrega parrafo con soporte para **negrita** inline."""
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # parsear **negrita**
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            p.add_run(part)
    return p


def add_table_from_md(doc, header_line, separator_line, rows_lines):
    """Construye tabla Word desde lineas markdown."""
    headers = [c.strip() for c in header_line.strip('|').split('|')]
    n_cols  = len(headers)
    n_rows  = len(rows_lines) + 1

    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.style = 'Table Grid'

    # encabezado
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(10)

    # filas
    for i, row_line in enumerate(rows_lines, 1):
        cols = [c.strip() for c in row_line.strip('|').split('|')]
        for j, val in enumerate(cols[:n_cols]):
            # limpiar markdown inline (**texto**)
            val_clean = re.sub(r'\*\*([^*]+)\*\*', r'\1', val)
            val_clean = re.sub(r'\*([^*]+)\*', r'\1', val_clean)
            val_clean = re.sub(r'`([^`]+)`', r'\1', val_clean)
            cell = table.cell(i, j)
            cell.text = val_clean
            cell.paragraphs[0].runs[0].font.size = Pt(10)

    doc.add_paragraph()  # espacio tras tabla


# ── parsear el markdown ───────────────────────────────────────────────────────
lines = MD_PATH.read_text(encoding='utf-8').splitlines()

i = 0
table_buffer = []
in_table = False

while i < len(lines):
    line = lines[i]

    # Saltar lineas de separador de tablas puras
    if re.match(r'^\|[-| :]+\|$', line):
        i += 1
        continue

    # Detectar inicio de tabla
    if line.startswith('|') and '|' in line[1:]:
        # recolectar tabla
        header_line = line
        i += 1
        sep_line = lines[i] if i < len(lines) else ''
        i += 1
        row_lines = []
        while i < len(lines) and lines[i].startswith('|'):
            if not re.match(r'^\|[-| :]+\|$', lines[i]):
                row_lines.append(lines[i])
            i += 1
        add_table_from_md(doc, header_line, sep_line, row_lines)
        continue

    # Encabezados
    if line.startswith('# ') and not line.startswith('## '):
        doc.add_heading(line[2:], level=1)
    elif line.startswith('## ') and not line.startswith('### '):
        doc.add_heading(line[3:], level=2)
    elif line.startswith('### '):
        doc.add_heading(line[4:], level=3)

    # Separador horizontal
    elif line.startswith('---'):
        doc.add_paragraph()

    # Imagen: ![alt](ruta)
    elif re.match(r'^!\[[^\]]*\]\([^)]+\)$', line.strip()):
        m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)$', line.strip())
        alt, ruta = m.group(1), m.group(2)
        ruta_abs = (BASE_DIR / ruta).resolve()
        if ruta_abs.exists():
            doc.add_picture(str(ruta_abs), width=Inches(6.3))
            last_p = doc.paragraphs[-1]
            last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if alt:
                cap = doc.add_paragraph(style='Normal')
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cap.add_run(alt)
                run.italic = True
                run.font.size = Pt(9)
        else:
            print(f'AVISO: imagen no encontrada, se omite: {ruta_abs}')

    # Blockquote (notas pendientes)
    elif line.startswith('> '):
        text = line[2:]
        text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
        p = doc.add_paragraph(style='Normal')
        p.paragraph_format.left_indent = Inches(0.4)
        run = p.add_run(text)
        run.italic = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(100, 100, 100)

    # Lista con viñeta
    elif line.startswith('- '):
        text = line[2:]
        p = doc.add_paragraph(style='List Bullet')
        parts = re.split(r'(\*\*[^*]+\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            else:
                # limpiar backticks
                sub_parts = re.split(r'(`[^`]+`)', part)
                for sp in sub_parts:
                    if sp.startswith('`') and sp.endswith('`'):
                        r = p.add_run(sp[1:-1])
                        r.font.name = 'Courier New'
                        r.font.size = Pt(10)
                    else:
                        p.add_run(sp)

    # Linea vacia
    elif line.strip() == '':
        pass  # saltar

    # Parrafo normal
    else:
        # limpiar backticks para codigos inline
        text = line
        p = doc.add_paragraph(style='Normal')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        # parsear **negrita** y `codigo`
        parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('`') and part.endswith('`'):
                run = p.add_run(part[1:-1])
                run.font.name = 'Courier New'
                run.font.size = Pt(10)
            elif part.startswith('*') and part.endswith('*'):
                run = p.add_run(part[1:-1])
                run.italic = True
            else:
                p.add_run(part)

    i += 1

doc.save(OUT_PATH)
print(f'Listo: {OUT_PATH}')
